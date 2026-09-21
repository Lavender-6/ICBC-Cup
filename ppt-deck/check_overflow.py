# -*- coding: utf-8 -*-
"""用真实字体度量估算每个文本框所需高度，找出可能溢出的框。"""
import sys
from pptx import Presentation
from pptx.util import Pt
from PIL import ImageFont

PPTX = "artifacts/presentation.pptx"
EMU_PT = 12700.0
REG = "C:/Windows/Fonts/msyh.ttc"
BOLD = "C:/Windows/Fonts/msyhbd.ttc"
_cache = {}


def font(size_pt, bold=False):
    key = (round(size_pt, 1), bold)
    if key not in _cache:
        px = int(round(size_pt * 96.0 / 72.0))
        _cache[key] = ImageFont.truetype(BOLD if bold else REG, px)
    return _cache[key]


def is_cjk(ch):
    o = ord(ch)
    return (0x4E00 <= o <= 0x9FFF) or (0x3000 <= o <= 0x303F) or \
           (0xFF00 <= o <= 0xFFEF) or (0x3400 <= o <= 0x4DBF)


def wrap(text, size_pt, bold, width_pt):
    """返回行数。CJK 可任意断行，拉丁按词断行。"""
    f = font(size_pt, bold)
    wpx = width_pt * 96.0 / 72.0
    if wpx <= 0:
        return 1
    # 先按显式换行拆段
    segs = text.split("\n")
    total = 0
    for seg in segs:
        if not seg:
            total += 1
            continue
        # 切成不可断单元
        units, buf = [], ""
        for ch in seg:
            if is_cjk(ch):
                if buf:
                    units.append(buf)
                    buf = ""
                units.append(ch)
            elif ch == " ":
                buf += ch
                units.append(buf)
                buf = ""
            else:
                buf += ch
        if buf:
            units.append(buf)
        cur = 0.0
        lines = 1
        for u in units:
            uw = f.getlength(u)
            if cur + uw > wpx and cur > 0:
                lines += 1
                cur = 0.0
                if u == " ":
                    continue
            cur += uw
        total += lines
    return total


def main():
    prs = Presentation(PPTX)
    report = []
    for idx, slide in enumerate(prs.slides, 1):
        for sh in slide.shapes:
            if not sh.has_text_frame:
                continue
            tf = sh.text_frame
            if not tf.text.strip():
                continue
            bw = (sh.width or 0) / EMU_PT
            bh = (sh.height or 0) / EMU_PT
            need = 0.0
            maxpt = 0.0
            for p in tf.paragraphs:
                txt = "".join(r.text for r in p.runs)
                if not txt:
                    continue
                pt = 0.0
                bold = False
                for r in p.runs:
                    if r.font.size:
                        pt = max(pt, r.font.size.pt)
                    bold = bold or bool(r.font.bold)
                if pt == 0:
                    pt = 18.0
                maxpt = max(maxpt, pt)
                ls = p.line_spacing
                if isinstance(ls, float) or isinstance(ls, int):
                    lsm = float(ls)
                else:
                    lsm = 1.2
                if lsm <= 0:
                    lsm = 1.2
                n = wrap(txt, pt, bold, bw)
                need += n * pt * lsm
            if need > bh * 1.02:
                report.append((idx, round(need - bh, 1), round(bh, 1), round(maxpt, 1),
                               tf.text.replace("\n", " / ")[:46]))
    report.sort(key=lambda r: -r[1])
    print("可能溢出的文本框：%d 个" % len(report))
    for r in report[:60]:
        print("P%02d  超出%6.1fpt  框高%5.1fpt  最大字号%4.1f | %s" % r)


if __name__ == "__main__":
    sys.stdout.reconfigure(encoding="utf-8")
    main()

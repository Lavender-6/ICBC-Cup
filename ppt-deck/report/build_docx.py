# -*- coding: utf-8 -*-
"""按挑战杯报告书格式生成「工银科创桥」项目报告书 .docx"""
import os
import re
import sys

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

sys.stdout.reconfigure(encoding="utf-8")

BASE = r"C:\Users\lenovo\Desktop\工行杯\ppt-deck\report"
OUT = r"C:\Users\lenovo\Desktop\工行杯\工银科创桥——硬科技企业里程碑式投贷联动平台-项目报告书.docx"

RED = RGBColor(0xC7, 0x00, 0x0B)
DARK = RGBColor(0x0A, 0x17, 0x30)
GREY = RGBColor(0x55, 0x55, 0x55)

CN_BODY = "宋体"
CN_HEAD = "黑体"
EN = "Times New Roman"


# ---------------- 基础工具 ----------------
def set_run(run, size=12, bold=False, cn=CN_BODY, en=EN, color=None, italic=False):
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.italic = italic
    run.font.name = en
    rPr = run._element.get_or_add_rPr()
    rFonts = rPr.find(qn("w:rFonts"))
    if rFonts is None:
        rFonts = OxmlElement("w:rFonts")
        rPr.insert(0, rFonts)
    rFonts.set(qn("w:ascii"), en)
    rFonts.set(qn("w:hAnsi"), en)
    rFonts.set(qn("w:eastAsia"), cn)
    if color is not None:
        run.font.color.rgb = color


def para(doc, text="", size=12, bold=False, cn=CN_BODY, en=EN, color=None,
         align="left", indent=True, space_before=0, space_after=6,
         line=1.5, left_indent=0):
    p = doc.add_paragraph()
    p.alignment = {"left": WD_ALIGN_PARAGRAPH.LEFT,
                   "center": WD_ALIGN_PARAGRAPH.CENTER,
                   "right": WD_ALIGN_PARAGRAPH.RIGHT,
                   "just": WD_ALIGN_PARAGRAPH.JUSTIFY}[align]
    pf = p.paragraph_format
    pf.space_before = Pt(space_before)
    pf.space_after = Pt(space_after)
    pf.line_spacing = line
    if left_indent:
        pf.left_indent = Pt(left_indent)
    if indent and not left_indent:
        pf.first_line_indent = Pt(size * 2)
    if text:
        r = p.add_run(text)
        set_run(r, size=size, bold=bold, cn=cn, en=en, color=color)
    return p


def heading(doc, text, level=1):
    if level == 1:
        p = para(doc, text, size=15, bold=True, cn=CN_HEAD, en=CN_HEAD,
                 color=DARK, align="left", indent=False,
                 space_before=16, space_after=10, line=1.4)
    elif level == 2:
        p = para(doc, text, size=13, bold=True, cn=CN_HEAD, en=CN_HEAD,
                 color=DARK, align="left", indent=False,
                 space_before=10, space_after=6, line=1.4)
    else:
        p = para(doc, text, size=12, bold=True, cn=CN_HEAD, en=CN_HEAD,
                 align="left", indent=False,
                 space_before=8, space_after=4, line=1.4)
    # 挂上 Word 大纲级别，便于导航窗格与目录域识别
    pPr = p._p.get_or_add_pPr()
    outline = OxmlElement("w:outlineLvl")
    outline.set(qn("w:val"), str(level - 1))
    pPr.append(outline)
    return p


def page_break(doc):
    doc.add_paragraph().add_run().add_break(WD_BREAK.PAGE)


def add_toc_field(doc):
    p = doc.add_paragraph()
    r = p.add_run()
    fld = OxmlElement("w:fldChar")
    fld.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = r'TOC \o "1-2" \h \z \u'
    sep = OxmlElement("w:fldChar")
    sep.set(qn("w:fldCharType"), "separate")
    txt = OxmlElement("w:t")
    txt.text = "【请在 Word 中右键“更新域”生成目录】"
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    for e in (fld, instr, sep, txt, end):
        r._r.append(e)
    set_run(r, size=11, color=GREY)


def add_page_number_footer(section):
    p = section.footer.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run()
    for tag, attr in (("w:fldChar", "begin"), ("w:instrText", None),
                      ("w:fldChar", "separate"), ("w:t", None), ("w:fldChar", "end")):
        if tag == "w:instrText":
            e = OxmlElement(tag)
            e.set(qn("xml:space"), "preserve")
            e.text = "PAGE"
        elif tag == "w:t":
            e = OxmlElement(tag)
            e.text = "1"
        else:
            e = OxmlElement(tag)
            e.set(qn("w:fldCharType"),
                  "begin" if attr == "begin" else ("separate" if attr == "separate" else "end"))
        r._r.append(e)
    set_run(r, size=10, color=GREY)


def clean_inline(s):
    s = s.replace("**", "").replace("~~", "")
    s = re.sub(r"(?<!\*)\*(?!\*)", "", s)
    s = re.sub(r"`([^`]*)`", r"\1", s)
    s = re.sub(r"\[([^\]]*)\]\([^)]*\)", r"\1", s)
    return s.strip()


def add_table(doc, rows):
    if not rows:
        return
    ncol = max(len(r) for r in rows)
    t = doc.add_table(rows=len(rows), cols=ncol)
    t.style = "Table Grid"
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    for i, row in enumerate(rows):
        for j in range(ncol):
            cell = t.cell(i, j)
            cell.text = ""
            p = cell.paragraphs[0]
            pf = p.paragraph_format
            pf.space_before = Pt(2)
            pf.space_after = Pt(2)
            pf.line_spacing = 1.25
            txt = clean_inline(row[j]) if j < len(row) else ""
            r = p.add_run(txt)
            set_run(r, size=10.5, bold=(i == 0), cn=CN_HEAD if i == 0 else CN_BODY,
                    en=CN_HEAD if i == 0 else EN)
    doc.add_paragraph().paragraph_format.space_after = Pt(4)


def render_markdown(doc, md_text):
    lines = md_text.split("\n")
    i = 0
    while i < len(lines):
        raw = lines[i].rstrip()
        line = raw.strip()

        if not line:
            i += 1
            continue

        # 表格
        if line.startswith("|"):
            block = []
            while i < len(lines) and lines[i].strip().startswith("|"):
                block.append(lines[i].strip())
                i += 1
            rows = []
            for r in block:
                cells = [c.strip() for c in r.strip("|").split("|")]
                if all(re.fullmatch(r":?-{2,}:?", c) for c in cells if c):
                    continue
                rows.append(cells)
            add_table(doc, rows)
            continue

        # 分隔线
        if re.fullmatch(r"-{3,}", line):
            i += 1
            continue

        # 标题
        m = re.match(r"^(#{1,6})\s+(.*)$", line)
        if m:
            heading(doc, clean_inline(m.group(2)), level=len(m.group(1)))
            i += 1
            continue

        # 引用
        if line.startswith(">"):
            txt = clean_inline(line.lstrip("> ").strip())
            para(doc, txt, size=11, color=GREY, indent=False,
                 left_indent=24, space_after=6, line=1.4)
            i += 1
            continue

        # 列表
        m = re.match(r"^([-*])\s+(.*)$", line)
        if m:
            para(doc, "· " + clean_inline(m.group(2)), size=12, indent=False,
                 left_indent=24, space_after=3, line=1.5)
            i += 1
            continue
        m = re.match(r"^(\d+)\.\s+(.*)$", line)
        if m:
            para(doc, m.group(1) + ". " + clean_inline(m.group(2)), size=12,
                 indent=False, left_indent=24, space_after=3, line=1.5)
            i += 1
            continue

        # 正文
        para(doc, clean_inline(line), size=12, align="just", indent=True,
             space_after=6, line=1.5)
        i += 1


# ---------------- 封面 ----------------
def build_cover(doc):
    for _ in range(3):
        para(doc, "", size=12, indent=False, space_after=0)
    para(doc, "第十七届「工行杯」全国大学生金融科技创新大赛", size=16, bold=True,
         cn=CN_HEAD, en=CN_HEAD, color=RED, align="center", indent=False, space_after=8)
    para(doc, "科 技 金 融 方 向 · 项 目 报 告 书", size=12, color=GREY,
         align="center", indent=False, space_after=30)
    para(doc, "工银科创桥", size=36, bold=True, cn=CN_HEAD, en=CN_HEAD,
         color=DARK, align="center", indent=False, space_after=10)
    para(doc, "硬科技企业里程碑式投贷联动平台", size=18, bold=True, cn=CN_HEAD,
         en=CN_HEAD, color=RED, align="center", indent=False, space_after=20)
    para(doc, "—— 让研发的每一步，都成为可计量的信用 ——", size=13,
         align="center", indent=False, space_after=40)

    info = [
        ("项目名称", "工银科创桥——硬科技企业里程碑式投贷联动平台"),
        ("参赛方向", "科技金融"),
        ("参赛高校", "【待填写】"),
        ("团队名称", "【待填写】"),
        ("团队成员", "【待填写】"),
        ("指导教师", "【待填写】"),
        ("完成日期", "2026 年 9 月"),
    ]
    t = doc.add_table(rows=len(info), cols=2)
    t.style = "Table Grid"
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    for i, (k, v) in enumerate(info):
        for j, txt in enumerate((k, v)):
            cell = t.cell(i, j)
            cell.text = ""
            p = cell.paragraphs[0]
            p.alignment = WD_ALIGN_PARAGRAPH.RIGHT if j == 0 else WD_ALIGN_PARAGRAPH.LEFT
            p.paragraph_format.space_before = Pt(3)
            p.paragraph_format.space_after = Pt(3)
            r = p.add_run(txt)
            set_run(r, size=11, bold=(j == 0), cn=CN_HEAD if j == 0 else CN_BODY,
                    en=CN_HEAD if j == 0 else EN)
    page_break(doc)


# ---------------- 主流程 ----------------
def read(name):
    with open(os.path.join(BASE, name), encoding="utf-8") as f:
        return f.read()


def main():
    doc = Document()

    sec = doc.sections[0]
    sec.page_width = Cm(21.0)
    sec.page_height = Cm(29.7)
    sec.left_margin = Cm(3.0)
    sec.right_margin = Cm(3.0)
    sec.top_margin = Cm(2.5)
    sec.bottom_margin = Cm(2.5)

    normal = doc.styles["Normal"]
    normal.font.name = EN
    normal.font.size = Pt(12)
    normal.element.rPr.rFonts.set(qn("w:eastAsia"), CN_BODY)

    add_page_number_footer(sec)

    build_cover(doc)

    render_markdown(doc, read("01_abstract.md"))
    page_break(doc)

    para(doc, "目  录", size=16, bold=True, cn=CN_HEAD, en=CN_HEAD,
         color=DARK, align="center", indent=False, space_after=14)
    add_toc_field(doc)
    para(doc, "（提示：在 Word 中打开后，右键目录区域选择“更新域”，即可自动生成页码）",
         size=10, color=GREY, align="center", indent=False, space_after=6)
    page_break(doc)

    for name in ("02_ch1_2.md", "03_ch3_4_5.md", "04_ch6_7_8.md", "05_ch9_10.md"):
        render_markdown(doc, read(name))

    doc.save(OUT)
    print("已生成：", OUT)


if __name__ == "__main__":
    main()
